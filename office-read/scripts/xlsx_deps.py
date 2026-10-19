#!/usr/bin/env python3
"""Trace formula dependencies in a workbook: what feeds a cell, or what a cell feeds.

Usage:
    xlsx_deps.py book.xlsx --cell "メイン!D10"                 precedents: the cells D10 is computed from
    xlsx_deps.py book.xlsx --cell "メイン!B3" --dependents     dependents: every formula that uses B3
    xlsx_deps.py book.xlsx --cell "メイン!D10" --depth 2
    xlsx_deps.py book.xlsx --all                              every formula cell with its direct references

Each node shows the formula and its cached value, so a wrong total can be followed back to the
input that caused it without reading every formula by hand. Values marked [uncached] come from a
workbook that was never recalculated; the formula is real, the number is unknown.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")  # Windows console defaults to cp932
except AttributeError:
    pass

RANGE_CAP = 2000          # cells; a bigger range is kept as one node instead of being expanded
CELL_RE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)$")
RANGE_RE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)$")
COLS_RE = re.compile(r"^\$?([A-Z]{1,3}):\$?([A-Z]{1,3})$")
ROWS_RE = re.compile(r"^\$?(\d+):\$?(\d+)$")


def col_num(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def col_letter(idx: int) -> str:
    s = ""
    while idx:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


def split_ref(ref: str, default_sheet: str) -> tuple[str, str]:
    """'Sheet'!A1 / Sheet!A1 / A1 -> (sheet, local)."""
    if "!" in ref:
        sheet, local = ref.rsplit("!", 1)
        sheet = sheet.strip("'").replace("''", "'")
        return sheet, local.upper()
    return default_sheet, ref.upper()


class Graph:
    def __init__(self, path: Path):
        import openpyxl
        self.wb = openpyxl.load_workbook(path, data_only=False)
        self.vals = openpyxl.load_workbook(path, data_only=True)
        self.formulas: dict[tuple[str, str], str] = {}
        self.precedents: dict[tuple[str, str], list[str]] = {}      # (sheet, cell) -> ["Sheet!A1", "Sheet!A1:B9"]
        self.rev: dict[tuple[str, str], set[tuple[str, str]]] = defaultdict(set)     # cell -> formula cells using it
        self.rev_ranges: list[tuple[str, int, int, int, int, tuple[str, str]]] = []    # big ranges, checked on query
        self.external: set[str] = set()
        self._build()

    # -- parsing ---------------------------------------------------------------------------

    def _expand(self, sheet: str, local: str) -> list[tuple[str, str]] | None:
        """Cells of a local reference, or None when it is too large to expand."""
        m = CELL_RE.match(local)
        if m:
            return [(sheet, f"{m[1]}{m[2]}")]
        m = RANGE_RE.match(local)
        if m:
            c1, r1, c2, r2 = col_num(m[1]), int(m[2]), col_num(m[3]), int(m[4])
            if (c2 - c1 + 1) * (r2 - r1 + 1) > RANGE_CAP:
                return None
            return [(sheet, f"{col_letter(c)}{r}") for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        return None     # whole columns/rows: never expanded

    def _bounds(self, local: str) -> tuple[int, int, int, int] | None:
        m = RANGE_RE.match(local)
        if m:
            return col_num(m[1]), int(m[2]), col_num(m[3]), int(m[4])
        m = COLS_RE.match(local)
        if m:
            return col_num(m[1]), 1, col_num(m[2]), 1_048_576
        m = ROWS_RE.match(local)
        if m:
            return 1, int(m[1]), 16_384, int(m[2])
        return None

    def _refs_in(self, formula: str, sheet: str) -> list[str]:
        from openpyxl.formula import Tokenizer
        out: list[str] = []
        try:
            tokens = Tokenizer(formula).items
        except Exception:
            return out
        for t in tokens:
            if t.type != "OPERAND" or t.subtype != "RANGE":
                continue
            v = t.value
            if v.startswith("[") or "[" in v.split("!")[0]:
                self.external.add(v)
                continue
            sh, local = split_ref(v, sheet)
            if CELL_RE.match(local) or RANGE_RE.match(local) or COLS_RE.match(local) or ROWS_RE.match(local):
                out.append(f"{sh}!{local.replace('$', '')}")
            else:
                # defined name
                dn = self.wb.defined_names.get(local) if hasattr(self.wb.defined_names, "get") else None
                if dn is not None:
                    for dsheet, dref in dn.destinations:
                        out.append(f"{dsheet}!{dref.replace('$', '').upper()}")
                else:
                    out.append(f"{sh}!{local}")
        return out

    def _build(self) -> None:
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        key = (ws.title, c.coordinate)
                        self.formulas[key] = c.value
                        refs = self._refs_in(c.value, ws.title)
                        self.precedents[key] = refs
                        for ref in refs:
                            sh, local = split_ref(ref, ws.title)
                            cells = self._expand(sh, local)
                            if cells is None:
                                b = self._bounds(local)
                                if b:
                                    self.rev_ranges.append((sh, *b, key))
                            else:
                                for cell in cells:
                                    self.rev[cell].add(key)

    # -- queries ---------------------------------------------------------------------------

    def value(self, sheet: str, cell: str) -> str:
        if sheet not in self.vals.sheetnames:
            return "(sheet not found)"
        f = self.formulas.get((sheet, cell))
        v = self.vals[sheet][cell].value
        if f is not None:
            return f"{f}  → {v}" if v is not None else f"{f}  [uncached]"
        return repr(v) if isinstance(v, str) else str(v)

    def dependents_of(self, sheet: str, cell: str) -> set[tuple[str, str]]:
        found = set(self.rev.get((sheet, cell), ()))
        m = CELL_RE.match(cell)
        if m:
            c, r = col_num(m[1]), int(m[2])
            for sh, c1, r1, c2, r2, key in self.rev_ranges:
                if sh == sheet and c1 <= c <= c2 and r1 <= r <= r2:
                    found.add(key)
        return found

    def tree_precedents(self, sheet: str, cell: str, depth: int) -> list[str]:
        lines: list[str] = []
        seen: set[tuple[str, str]] = set()

        def walk(sh: str, local: str, level: int, path: set):
            key = (sh, local)
            indent = "  " * level
            lines.append(f"{indent}- {sh}!{local} = {self.value(sh, local)}")
            if key in path:
                lines.append(f"{indent}  - (circular reference)")
                return
            if key not in self.precedents or level >= depth:
                if key in self.precedents and level >= depth:
                    lines.append(f"{indent}  - … (depth limit; rerun with --depth {depth + 2})")
                return
            if key in seen:
                lines.append(f"{indent}  - (shown above)")
                return
            seen.add(key)
            for ref in self.precedents[key]:
                rsh, rlocal = split_ref(ref, sh)
                cells = self._expand(rsh, rlocal)
                if cells is None or len(cells) > 1:
                    n = len(cells) if cells else "many"
                    sub_formulas = [c for c in (cells or []) if c in self.formulas]
                    lines.append(f"{indent}  - {rsh}!{rlocal} (range, {n} cells, {len(sub_formulas)} with formulas)")
                    for c in sub_formulas[:8]:
                        walk(c[0], c[1], level + 2, path | {key})
                    if len(sub_formulas) > 8:
                        lines.append(f"{indent}    - … {len(sub_formulas) - 8} more formula cells in this range")
                else:
                    walk(cells[0][0], cells[0][1], level + 1, path | {key})

        walk(sheet, cell, 0, set())
        return lines

    def tree_dependents(self, sheet: str, cell: str, depth: int) -> list[str]:
        lines: list[str] = []
        seen: set[tuple[str, str]] = set()

        def walk(sh: str, local: str, level: int, path: set):
            key = (sh, local)
            indent = "  " * level
            lines.append(f"{indent}- {sh}!{local} = {self.value(sh, local)}")
            if key in path:
                lines.append(f"{indent}  - (circular reference)")
                return
            deps = sorted(self.dependents_of(sh, local))
            if not deps or level >= depth:
                if deps and level >= depth:
                    lines.append(f"{indent}  - … {len(deps)} dependent(s) beyond depth limit; rerun with --depth {depth + 2}")
                return
            if key in seen:
                lines.append(f"{indent}  - (shown above)")
                return
            seen.add(key)
            for d in deps:
                walk(d[0], d[1], level + 1, path | {key})

        walk(sheet, cell, 0, set())
        return lines


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--cell", help="Sheet!A1 (sheet optional when the workbook has one sheet)")
    ap.add_argument("--dependents", action="store_true", help="what uses the cell, instead of what feeds it")
    ap.add_argument("--depth", type=int, default=5)
    ap.add_argument("--all", action="store_true", help="list every formula cell with its direct references")
    a = ap.parse_args(argv)
    if not a.file.exists():
        sys.exit(f"[xlsx_deps] file not found: {a.file}")
    if not a.cell and not a.all:
        sys.exit("[xlsx_deps] pass --cell Sheet!A1 or --all")

    g = Graph(a.file)
    print(f"# {a.file.name}")
    print()
    if a.all:
        print(f"{len(g.formulas)} formula cells")
        print()
        for (sh, cell), f in sorted(g.formulas.items()):
            refs = ", ".join(g.precedents[(sh, cell)]) or "(no cell references)"
            print(f"- {sh}!{cell} = {g.value(sh, cell)}  ← {refs}")
    else:
        sheet, local = split_ref(a.cell, g.wb.sheetnames[0] if len(g.wb.sheetnames) == 1 else "")
        if not sheet:
            sys.exit(f"[xlsx_deps] workbook has several sheets; write the cell as Sheet!A1. Sheets: {g.wb.sheetnames}")
        if sheet not in g.wb.sheetnames:
            sys.exit(f"[xlsx_deps] sheet '{sheet}' not found. Sheets: {g.wb.sheetnames}")
        local = local.replace("$", "")
        if a.dependents:
            print(f"Dependents of {sheet}!{local} (formulas that use it), depth {a.depth}:")
            print()
            print("\n".join(g.tree_dependents(sheet, local, a.depth)))
        else:
            print(f"Precedents of {sheet}!{local} (what it is computed from), depth {a.depth}:")
            print()
            print("\n".join(g.tree_precedents(sheet, local, a.depth)))
    if g.external:
        print()
        print("> External workbook references (not followed): " + ", ".join(sorted(g.external)[:10]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
