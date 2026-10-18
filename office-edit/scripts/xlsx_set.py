#!/usr/bin/env python3
"""Set cell values / formulas in an existing workbook without writing throwaway code.

Keeps a backup, preserves VBA in .xlsm, and prints what changed. For anything beyond
single-cell writes (inserting rows, formatting, charts) write openpyxl code directly;
this is the quick path for "put X in Sheet!A1" style edits.

Usage:
    xlsx_set.py book.xlsx --set "売上!D6==SUM(D2:D5)" --set "売上!A7=合計"
    xlsx_set.py book.xlsx --json edits.json          # {"売上!D6": "=SUM(D2:D5)", "Sheet2!B2": 42}
    xlsx_set.py book.xlsx --set "Sheet1!A1=hi" --out copy.xlsx   # leave original untouched
    xlsx_set.py book.xlsx --set ... --dry-run

Value parsing: "=..." is a formula; integers/floats become numbers; "true"/"false" become
booleans; anything else is text. Prefix with a single quote to force text ("'0123").
"""

from __future__ import annotations

import argparse
from copy import copy
import json
import re
import shutil
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")  # Windows console defaults to cp932
except AttributeError:
    pass


def parse_value(raw):
    if isinstance(raw, (int, float, bool)) or raw is None:
        return raw
    s = str(raw)
    if s.startswith("'"):
        return s[1:]
    if s.startswith("="):
        return s
    if s.lower() in ("true", "false"):
        return s.lower() == "true"
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d+\.\d*|-?\.\d+", s):
        return float(s)
    return s


def split_ref(ref: str) -> tuple[str | None, str]:
    if "!" in ref:
        sheet, cell = ref.rsplit("!", 1)
        return sheet.strip("'"), cell.upper()
    return None, ref.upper()


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--set", action="append", default=[], metavar="SHEET!CELL=VALUE")
    ap.add_argument("--json", type=Path, help="JSON object of ref -> value")
    ap.add_argument("--out", type=Path, help="write here instead of overwriting (no backup then)")
    ap.add_argument("--no-backup", action="store_true")
    ap.add_argument("--no-copy-style", action="store_true",
                    help="don't copy number format/font/border from the cell above into an unstyled target cell")
    ap.add_argument("--dry-run", action="store_true")
    opts = ap.parse_args(argv)

    edits: dict[str, object] = {}
    for s in opts.set:
        if "=" not in s:
            sys.exit(f"[xlsx_set] bad --set (expected REF=VALUE): {s}")
        ref, val = s.split("=", 1)
        edits[ref] = val
    if opts.json:
        edits.update(json.loads(opts.json.read_text(encoding="utf-8")))
    if not edits:
        sys.exit("[xlsx_set] nothing to do; pass --set or --json")

    import openpyxl
    keep_vba = opts.file.suffix.lower() in (".xlsm", ".xltm")
    wb = openpyxl.load_workbook(opts.file, keep_vba=keep_vba)

    changes = []
    for ref, raw in edits.items():
        sheet, cell = split_ref(ref)
        ws = wb[sheet] if sheet else wb.active
        if sheet and sheet not in wb.sheetnames:
            sys.exit(f"[xlsx_set] sheet not found: {sheet}. Sheets: {wb.sheetnames}")
        new = parse_value(raw)
        target_cell = ws[cell]
        old = target_cell.value
        styled = ""
        # New rows are usually appended under formatted ones; an unstyled cell next to a styled
        # column looks wrong in Excel, so inherit from the cell above unless told not to.
        above = ws.cell(row=target_cell.row - 1, column=target_cell.column) if target_cell.row > 1 else None
        if (not opts.no_copy_style and above is not None and above.has_style and not target_cell.has_style):
            if not opts.dry_run:
                target_cell._style = copy(above._style)
            styled = f"  (style copied from {above.coordinate})"
        changes.append((ws.title, cell, old, new, styled))
        if not opts.dry_run:
            target_cell.value = new

    for sheet, cell, old, new, styled in changes:
        print(f"  {sheet}!{cell}: {old!r} -> {new!r}{styled}")
    if opts.dry_run:
        print("[xlsx_set] dry run, nothing written")
        return 0

    target = opts.out or opts.file
    if target == opts.file and not opts.no_backup:
        bak = opts.file.with_suffix(opts.file.suffix + ".bak")
        shutil.copyfile(opts.file, bak)
        print(f"[xlsx_set] backup: {bak}")
    wb.save(target)
    print(f"[xlsx_set] wrote {target} ({len(changes)} cells). "
          "Formulas have no cached values yet; run recalc.py if anyone will read values programmatically.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
