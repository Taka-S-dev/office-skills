#!/usr/bin/env python3
"""Find-and-replace in a .pptx across slides, tables, grouped shapes and speaker notes.

Formatting is kept (replacement takes the first matched run's style). Use --dry-run to see
hit counts per pattern before writing.

Usage:
    pptx_replace.py deck.pptx --replace "FY2025=FY2026" --replace "山田=佐藤"
    pptx_replace.py deck.pptx --json map.json
    pptx_replace.py deck.pptx --replace "..." --slides 3,5-7      # restrict to these slides
    pptx_replace.py deck.pptx --replace "..." --no-notes
    pptx_replace.py deck.pptx --replace "..." --dry-run
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")  # Windows console defaults to cp932
except AttributeError:
    pass

sys.path.insert(0, str(Path(__file__).parent))
from _runs import find_in_runs, replace_in_runs  # noqa: E402


def parse_slides(spec: str | None) -> set[int] | None:
    if not spec:
        return None
    out: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            out.update(range(int(a), int(b) + 1))
        elif part:
            out.add(int(part))
    return out


def iter_text_frames(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_text_frames(sh.shapes)
            continue
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--replace", action="append", default=[], metavar="OLD=NEW")
    ap.add_argument("--json", type=Path)
    ap.add_argument("--regex", action="store_true")
    ap.add_argument("--slides", help="e.g. 2,4-6 (1-based); default all")
    ap.add_argument("--no-notes", action="store_true", help="skip speaker notes")
    ap.add_argument("--out", type=Path)
    ap.add_argument("--no-backup", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    opts = ap.parse_args(argv)

    pairs: dict[str, str] = {}
    for s in opts.replace:
        if "=" not in s:
            sys.exit(f"[pptx_replace] bad --replace (expected OLD=NEW): {s}")
        old, new = s.split("=", 1)
        pairs[old] = new
    if opts.json:
        pairs.update(json.loads(opts.json.read_text(encoding="utf-8")))
    if not pairs:
        sys.exit("[pptx_replace] nothing to do; pass --replace or --json")

    from pptx import Presentation
    prs = Presentation(opts.file)
    only = parse_slides(opts.slides)

    totals = {old: 0 for old in pairs}
    per_slide: dict[int, int] = {}
    for n, slide in enumerate(prs.slides, start=1):
        if only and n not in only:
            continue
        frames = list(iter_text_frames(slide.shapes))
        if not opts.no_notes and slide.has_notes_slide:
            frames.append(slide.notes_slide.notes_text_frame)
        for tf in frames:
            for p in tf.paragraphs:
                runs = list(p.runs)
                for old, new in pairs.items():
                    k = (find_in_runs(runs, old, opts.regex) if opts.dry_run
                         else replace_in_runs(runs, old, new, opts.regex))
                    totals[old] += k
                    per_slide[n] = per_slide.get(n, 0) + k

    for old, n in totals.items():
        flag = "" if n else "   <-- not found"
        print(f"  {n:4d}  {old!r} -> {pairs[old]!r}{flag}")
    if per_slide:
        print("  slides touched: " + ", ".join(f"{k}({v})" for k, v in sorted(per_slide.items()) if v))

    if opts.dry_run:
        print("[pptx_replace] dry run, nothing written")
        return 0 if all(totals.values()) else 1

    target = opts.out or opts.file
    if target == opts.file and not opts.no_backup:
        bak = opts.file.with_suffix(opts.file.suffix + ".bak")
        shutil.copyfile(opts.file, bak)
        print(f"[pptx_replace] backup: {bak}")
    prs.save(target)
    print(f"[pptx_replace] wrote {target}")
    if any(not n for n in totals.values()):
        print("[pptx_replace] some patterns were not found. Charts, SmartArt and text inside "
              "images are not searchable; check with office-read if unexpected.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
