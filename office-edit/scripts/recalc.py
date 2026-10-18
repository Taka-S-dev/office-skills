#!/usr/bin/env python3
"""Recalculate an .xlsx/.xlsm written by openpyxl and report formula errors.

openpyxl writes formulas but never computes them, so a freshly saved file has empty
cached values until something opens and recalculates it. This script does that and then
scans every sheet for #REF!, #DIV/0!, #NAME?, #VALUE!, #N/A so you can fix mistakes
before handing the file over.

Engines, tried in order unless --engine is given:
  excel   Excel via COM (Windows + Excel installed; no Python package needed)
  soffice LibreOffice headless (any OS; needs `soffice` on PATH or in the usual install dir)

Usage:
    recalc.py book.xlsx                  recalc in place, print error report
    recalc.py book.xlsx --engine soffice
    recalc.py book.xlsx --check-only     no recalc, just scan cached values for errors
"""

from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")  # Windows console defaults to cp932
except AttributeError:
    pass

ERRORS = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!")


def find_soffice() -> str | None:
    exe = shutil.which("soffice") or shutil.which("soffice.exe")
    if exe:
        return exe
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice",
              "/usr/bin/soffice", "/usr/lib/libreoffice/program/soffice"):
        if Path(p).exists():
            return p
    return None


def recalc_excel(path: Path) -> None:
    """Open in Excel through COM, full-recalculate, save. Runs via PowerShell so pywin32 is optional.

    The path travels as an environment variable, never as script text, so no quoting rules apply
    and a file name cannot inject PowerShell. AutomationSecurity=3 keeps the workbook's macros
    from running: an untrusted file should never get code execution just for being recalculated.
    """
    ps = """
$ErrorActionPreference = 'Stop'
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false
$xl.DisplayAlerts = $false
$xl.AutomationSecurity = 3
try {
  $wb = $xl.Workbooks.Open($env:OFFICE_SRC)
  $xl.CalculateFullRebuild()
  $wb.Save()
  $wb.Close($false)
} finally {
  $xl.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($xl) | Out-Null
}
"""
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                       capture_output=True, text=True,
                       env={**os.environ, "OFFICE_SRC": str(path.resolve())})
    if r.returncode != 0:
        raise RuntimeError(f"Excel COM failed:\n{r.stderr.strip()}")


def recalc_soffice(path: Path, exe: str) -> None:
    """Round-trip through LibreOffice. It recalculates on load and writes cached values on save."""
    with tempfile.TemporaryDirectory() as td:
        ext = path.suffix.lstrip(".")
        r = subprocess.run([exe, "--headless", "--calc", "--convert-to", ext, "--outdir", td, str(path)],
                           capture_output=True, text=True, timeout=300)
        out = Path(td) / path.name
        if r.returncode != 0 or not out.exists():
            raise RuntimeError(f"soffice failed:\n{r.stdout}\n{r.stderr}")
        shutil.copyfile(out, path)


def scan_errors(path: Path) -> list[tuple[str, str, str]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value in ERRORS:
                    found.append((ws.title, c.coordinate, c.value))
    return found


def count_uncached(path: Path) -> int:
    """Formula cells whose cached value is still None after recalc = engine did not run."""
    import openpyxl
    f = openpyxl.load_workbook(path, data_only=False)
    v = openpyxl.load_workbook(path, data_only=True)
    n = 0
    for ws in f.worksheets:
        vs = v[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and vs[c.coordinate].value is None:
                    n += 1
    return n


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--engine", choices=["excel", "soffice"], help="force an engine")
    ap.add_argument("--check-only", action="store_true", help="skip recalculation, only scan for errors")
    opts = ap.parse_args(argv)
    path = opts.file
    if not path.exists():
        sys.exit(f"[recalc] file not found: {path}")

    if not opts.check_only:
        engine = opts.engine
        soffice = find_soffice()
        if engine is None:
            engine = "excel" if os.name == "nt" else ("soffice" if soffice else None)
        if engine == "excel":
            try:
                recalc_excel(path)
            except RuntimeError as e:
                if soffice and not opts.engine:
                    print(f"[recalc] Excel unavailable ({e.splitlines()[0]}), falling back to LibreOffice")
                    engine = "soffice"
                else:
                    sys.exit(f"[recalc] {e}")
        if engine == "soffice":
            if not soffice:
                sys.exit("[recalc] no engine: Excel COM not usable and LibreOffice not found. "
                         "Open the file in Excel once to recalculate, or install LibreOffice.")
            recalc_soffice(path, soffice)
        if engine is None:
            sys.exit("[recalc] no recalculation engine available on this OS.")
        print(f"[recalc] recalculated with {engine}: {path}")

    uncached = count_uncached(path)
    if uncached:
        print(f"[recalc] WARNING: {uncached} formula cells still have no cached value.")
    errs = scan_errors(path)
    if errs:
        print(f"[recalc] {len(errs)} formula error(s):")
        for sheet, cell, val in errs[:50]:
            print(f"  {sheet}!{cell}: {val}")
        if len(errs) > 50:
            print(f"  ... and {len(errs) - 50} more")
        return 2
    print("[recalc] no formula errors found.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
