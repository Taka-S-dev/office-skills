#!/usr/bin/env python3
"""Find-and-replace in a .docx, keeping formatting, across body, tables, headers and footers.

Matches that span several runs are handled (see _runs.py). Use --dry-run first: it prints
how many hits each pattern has, so you know the text exists before touching the file.

Usage:
    docx_replace.py report.docx --replace "2025年度=2026年度" --replace "株式会社A=株式会社B"
    docx_replace.py report.docx --json map.json            # {"old": "new", ...}
    docx_replace.py report.docx --replace "..." --dry-run
    docx_replace.py report.docx --replace "第(\\d+)回=第\\1期" --regex
    docx_replace.py report.docx --replace "..." --out edited.docx
    docx_replace.py report.docx --replace "..." --scope body,tables   # leave header/footer alone

The report shows where hits are (body / tables / header / footer), so a dry run tells you
whether a pattern would also touch the header before you commit.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")  # Windows console defaults to cp932
except AttributeError:
    pass

sys.path.insert(0, str(Path(__file__).parent))
from _runs import find_in_runs, replace_in_runs  # noqa: E402


def paragraph_runs(paragraph):
    """All runs including those inside hyperlinks, in document order."""
    from docx.oxml.ns import qn
    from docx.text.run import Run
    return [Run(r, paragraph) for r in paragraph._p.iter(qn("w:r"))]


SCOPES = ("body", "tables", "header", "footer")


def iter_paragraphs(doc):
    """Yield (scope, paragraph) for body, tables (nested too), headers, footers."""
    def from_tables(tables, scope):
        for t in tables:
            for row in t.rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        yield scope, p
                    yield from from_tables(cell.tables, scope)

    for p in doc.paragraphs:
        yield "body", p
    yield from from_tables(doc.tables, "tables")
    for s in doc.sections:
        for scope, parts in (("header", (s.header, s.first_page_header, s.even_page_header)),
                             ("footer", (s.footer, s.first_page_footer, s.even_page_footer))):
            for part in parts:
                if part is not None and not part.is_linked_to_previous:
                    for p in part.paragraphs:
                        yield scope, p
                    yield from from_tables(part.tables, scope)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--replace", action="append", default=[], metavar="OLD=NEW")
    ap.add_argument("--json", type=Path, help="JSON object of old -> new")
    ap.add_argument("--regex", action="store_true", help="treat OLD as a regex, NEW may use \\1")
    ap.add_argument("--scope", default="body,tables,header,footer",
                    help="comma list of body,tables,header,footer (default: all)")
    ap.add_argument("--out", type=Path)
    ap.add_argument("--no-backup", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    opts = ap.parse_args(argv)
    scopes = {x.strip() for x in opts.scope.split(",") if x.strip()}
    bad = scopes - set(SCOPES)
    if bad:
        sys.exit(f"[docx_replace] unknown scope {sorted(bad)}; choose from {SCOPES}")

    pairs: dict[str, str] = {}
    for s in opts.replace:
        if "=" not in s:
            sys.exit(f"[docx_replace] bad --replace (expected OLD=NEW): {s}")
        old, new = s.split("=", 1)
        pairs[old] = new
    if opts.json:
        pairs.update(json.loads(opts.json.read_text(encoding="utf-8")))
    if not pairs:
        sys.exit("[docx_replace] nothing to do; pass --replace or --json")

    import docx
    doc = docx.Document(opts.file)
    paragraphs = list(iter_paragraphs(doc))  # (scope, paragraph)

    totals = {old: 0 for old in pairs}
    where = {old: {sc: 0 for sc in SCOPES} for old in pairs}   # hits per scope, counted everywhere
    for scope, p in paragraphs:
        runs = paragraph_runs(p)
        for old, new in pairs.items():
            if opts.dry_run or scope not in scopes:
                k = find_in_runs(runs, old, opts.regex)
            else:
                k = replace_in_runs(runs, old, new, opts.regex)
            where[old][scope] += k
            if scope in scopes:
                totals[old] += k

    for old, n in totals.items():
        parts = [f"{sc}={c}" + ("" if sc in scopes else "(skipped)") for sc, c in where[old].items() if c]
        if n:
            flag = ""
        elif any(where[old].values()):
            flag = "   <-- hits only outside --scope, nothing changed"
        else:
            flag = "   <-- not found"
        print(f"  {n:4d}  {old!r} -> {pairs[old]!r}  [{', '.join(parts) or 'no hits'}]{flag}")

    if opts.dry_run:
        print("[docx_replace] dry run, nothing written")
        return 0 if all(totals.values()) else 1

    target = opts.out or opts.file
    if target == opts.file and not opts.no_backup:
        bak = opts.file.with_suffix(opts.file.suffix + ".bak")
        shutil.copyfile(opts.file, bak)
        print(f"[docx_replace] backup: {bak}")
    doc.save(target)
    print(f"[docx_replace] wrote {target}")
    if any(not n for n in totals.values()):
        print("[docx_replace] some patterns were not found. Text boxes, footnotes and "
              "tracked-change deletions are not searched; check with office-read if unexpected.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
